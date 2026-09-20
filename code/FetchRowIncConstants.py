"""
Fetches the raw external data needed to construct Rest-of-World wages
(MakeRowIncConstants.do turns these into Wᵈ_ROW, Wᶠ_ROW): World Bank GDP and
labor force totals for the World and the United States, 2015, and IRS
Statistics of Income Table 2 (Individual Income Tax Returns With Form 2555,
foreign-earned income by country/region) for tax years 2011 and 2016, used to
interpolate a 2015 figure for Americans' foreign-earned income.

Run this to (re)generate data/WorldBankGdpLaborForce2015.csv and
data/IrsForeignEarnedIncome.csv from their live sources.
"""

import io
import csv

import requests
import openpyxl

DATA_DIR = "../../data"

# %% World Bank: GDP (current US$) and total labor force, World and US, 2015
WORLD_BANK_INDICATORS = {
    "NY.GDP.MKTP.CD": "GDP (current US$)",
    "SL.TLF.TOTL.IN": "Labor force, total",
}

def fetch_world_bank():
    rows = []
    for indicator in WORLD_BANK_INDICATORS:
        url = (
            f"https://api.worldbank.org/v2/country/WLD;USA/indicator/{indicator}"
            "?format=json&date=2015&per_page=20"
        )
        resp = requests.get(url, timeout=30)
        resp.raise_for_status()
        _, records = resp.json()
        for rec in records:
            rows.append({
                "country": rec["country"]["value"],
                "indicator": indicator,
                "year": rec["date"],
                "value": rec["value"],
            })
    with open(f"{DATA_DIR}/WorldBankGdpLaborForce2015.csv", "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["country", "indicator", "year", "value"])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {DATA_DIR}/WorldBankGdpLaborForce2015.csv ({len(rows)} rows)")


# %% IRS SOI Table 2: Total foreign-earned income, "All geographic areas" row
IRS_TABLE2_URLS = {
    2011: "https://www.irs.gov/pub/irs-soi/11in02ic.xlsx",
    2016: "https://www.irs.gov/pub/irs-soi/16in02ic.xlsx",
}

def fetch_irs_foreign_earned_income():
    rows = []
    for year, url in IRS_TABLE2_URLS.items():
        resp = requests.get(url, timeout=30)
        resp.raise_for_status()
        wb = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]

        total_row = None
        for row in ws.iter_rows(values_only=True):
            label = row[0]
            if isinstance(label, str) and "all geographic areas" in label.lower():
                total_row = row
                break
        if total_row is None:
            raise RuntimeError(f"Could not find the 'All geographic areas' row in {url}")

        filers = total_row[1]
        total_foreign_earned_income_usd = total_row[2] * 1000  # table is in thousands of dollars
        rows.append({
            "tax_year": year,
            "filers": filers,
            "total_foreign_earned_income_usd": total_foreign_earned_income_usd,
            "weighted_average_usd": total_foreign_earned_income_usd / filers,
        })

    with open(f"{DATA_DIR}/IrsForeignEarnedIncome.csv", "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "tax_year", "filers", "total_foreign_earned_income_usd", "weighted_average_usd"
        ])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {DATA_DIR}/IrsForeignEarnedIncome.csv ({len(rows)} rows)")


if __name__ == "__main__":
    fetch_world_bank()
    fetch_irs_foreign_earned_income()
